# -*- coding: utf-8 -*-
# Change log:
# 2026-09-15: Shared OKA Excel import used by manage.py and the systemadministrator upload page.

from django.db import transaction
from systemoversikt.models import ApplicationLog, InformasjonsKategori
from systemoversikt.oka_kode import (
	NIVAA_FUNKSJON,
	NIVAA_HOVED,
	NIVAA_UNDER,
	nivaa_from_kode,
	normalize_kode,
	parent_kode_from_kode,
)
import io
import re
import time
import zipfile


LOG_EVENT_TYPE = 'OKA-klassifikasjon import'
MAX_UPLOAD_BYTES = 15 * 1024 * 1024


class OkaImportError(ValueError):
	pass


def _xlsx_without_autofilter(source):
	"""Strip broken AutoFilter XML so openpyxl can read the OKA workbook."""
	if hasattr(source, 'read'):
		data = source.read()
		if hasattr(source, 'seek'):
			try:
				source.seek(0)
			except Exception:
				pass
		zip_source = io.BytesIO(data)
	else:
		zip_source = source
	src = zipfile.ZipFile(zip_source, 'r')
	buf = io.BytesIO()
	with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as dest:
		for item in src.infolist():
			payload = src.read(item.filename)
			if item.filename.startswith('xl/worksheets/') and item.filename.endswith('.xml'):
				payload = re.sub(br'<autoFilter[^>]*>.*?</autoFilter>', b'', payload, flags=re.DOTALL)
				payload = re.sub(br'<autoFilter[^/]*/>', b'', payload)
			dest.writestr(item, payload)
	src.close()
	buf.seek(0)
	return buf


def _cell_text(value):
	if value is None:
		return ''
	return str(value).strip()


def _header_map(header_row):
	mapping = {}
	for idx, raw in enumerate(header_row):
		if raw is None:
			continue
		key = re.sub(r'[^a-z0-9]+', '', str(raw).strip().lower())
		if not key:
			continue
		if key.startswith('kode'):
			mapping['kode'] = idx
		elif key.startswith('tittel'):
			mapping['tittel'] = idx
		elif key.startswith('herlegges'):
			mapping['her_legges'] = idx
		elif key.startswith('bkvurdering'):
			mapping['bk_vurdering'] = idx
		elif 'tidligerearkiv' in key:
			mapping['tidligere_arkivnokkel'] = idx
		elif 'aktueltlovverk' in key:
			mapping['aktuelt_lovverk'] = idx
		elif key.startswith('kommentar'):
			mapping['kommentar'] = idx
	return mapping


def _read_sheet_rows(ws):
	rows = list(ws.iter_rows(values_only=True))
	if not rows:
		return []
	header_map = _header_map(rows[0])
	if 'kode' not in header_map or 'tittel' not in header_map:
		return []
	parsed = []
	for row in rows[1:]:
		kode = _cell_text(row[header_map['kode']]) if len(row) > header_map['kode'] else ''
		tittel = _cell_text(row[header_map['tittel']]) if len(row) > header_map['tittel'] else ''
		if not kode or not tittel:
			continue
		nivaa = nivaa_from_kode(kode)
		if not nivaa:
			continue

		def col(name):
			idx = header_map.get(name)
			if idx is None or len(row) <= idx:
				return ''
			return _cell_text(row[idx])

		parsed.append({
			'kode': kode,
			'kode_normalisert': normalize_kode(kode),
			'tittel': tittel,
			'nivaa': nivaa,
			'parent_kode': parent_kode_from_kode(kode),
			'her_legges': col('her_legges'),
			'bk_vurdering': col('bk_vurdering'),
			'tidligere_arkivnokkel': col('tidligere_arkivnokkel'),
			'aktuelt_lovverk': col('aktuelt_lovverk'),
			'kommentar': col('kommentar'),
		})
	return parsed


def import_oka_workbook(source, filename=''):
	"""Import OKA classification from an xlsx path or uploaded file-like object."""
	import openpyxl

	ApplicationLog.objects.create(
		event_type=LOG_EVENT_TYPE,
		message='starter%s..' % ((' (%s)' % filename) if filename else ''),
	)
	runtime_t0 = time.time()

	buf = _xlsx_without_autofilter(source)
	wb = openpyxl.load_workbook(buf, data_only=True, read_only=True)

	rows_by_kode = {}
	order = 0
	for sheet_name in wb.sheetnames:
		if sheet_name.strip().lower() == 'veiledning':
			continue
		for row in _read_sheet_rows(wb[sheet_name]):
			order += 1
			row['rekkefolge'] = order
			rows_by_kode[row['kode_normalisert']] = row
	wb.close()

	if not rows_by_kode:
		raise OkaImportError('Ingen rader med kode og tittel ble funnet.')

	level_rank = {NIVAA_HOVED: 0, NIVAA_FUNKSJON: 1, NIVAA_UNDER: 2}
	sorted_rows = sorted(
		rows_by_kode.values(),
		key=lambda r: (level_rank.get(r['nivaa'], 9), r['rekkefolge']),
	)

	ant_nye = 0
	ant_oppdatert = 0
	ant_deaktivert = 0
	seen = set()

	with transaction.atomic():
		by_norm = {
			obj.kode_normalisert: obj
			for obj in InformasjonsKategori.objects.all()
		}
		for row in sorted_rows:
			seen.add(row['kode_normalisert'])
			obj = by_norm.get(row['kode_normalisert'])
			if obj is None:
				obj = InformasjonsKategori(kode_normalisert=row['kode_normalisert'])
				ant_nye += 1
			else:
				ant_oppdatert += 1
			obj.kode = row['kode']
			obj.tittel = row['tittel']
			obj.nivaa = row['nivaa']
			obj.her_legges = row['her_legges']
			obj.bk_vurdering = row['bk_vurdering']
			obj.tidligere_arkivnokkel = row['tidligere_arkivnokkel']
			obj.aktuelt_lovverk = row['aktuelt_lovverk']
			obj.kommentar = row['kommentar']
			obj.aktiv = True
			obj.rekkefolge = row['rekkefolge']
			obj.save()
			by_norm[row['kode_normalisert']] = obj

		for row in sorted_rows:
			obj = by_norm[row['kode_normalisert']]
			parent = None
			if row['parent_kode']:
				parent = by_norm.get(normalize_kode(row['parent_kode']))
			if obj.parent_id != (parent.pk if parent else None):
				obj.parent = parent
				obj.save(update_fields=['parent'])

		for obj in InformasjonsKategori.objects.exclude(kode_normalisert__in=seen).filter(aktiv=True):
			obj.aktiv = False
			obj.save(update_fields=['aktiv'])
			ant_deaktivert += 1

	runtime = '%.1f' % (time.time() - runtime_t0)
	totalt_aktive = InformasjonsKategori.objects.filter(aktiv=True).count()
	message = (
		'OKA-import ferdig. Nye: %s, oppdatert: %s, deaktivert: %s. '
		'Totalt aktive: %s. Tid %ss.'
		% (ant_nye, ant_oppdatert, ant_deaktivert, totalt_aktive, runtime)
	)
	ApplicationLog.objects.create(event_type=LOG_EVENT_TYPE, message=message)
	return {
		'ant_nye': ant_nye,
		'ant_oppdatert': ant_oppdatert,
		'ant_deaktivert': ant_deaktivert,
		'totalt_aktive': totalt_aktive,
		'message': message,
	}
